import pytest
from openpyxl import load_workbook
from cablecalc.io_excel import (COLUMNS, GUIDE, METHOD_GUIDE, METHODS, Project, read_input,
                                write_template)
from cablecalc.models import InputError

ROW = {"tag": "C1", "phase": "3ph", "voltage_v": 400, "design_current_a": 30, "load_kw": None,
       "power_factor": 0.8, "length_m": 50, "conductor": "Cu", "insulation": "PVC",
       "method": "C", "cable": "3C", "ambient_c": 40, "grouped_circuits": 1,
       "vd_limit_pct": 5, "fault_current_ka": 2, "fault_time_s": 0.1, "device_rating_a": None}


def filled(tmp_path, rows, code="IEC"):
    p = tmp_path / "in.xlsx"
    write_template(p)
    wb = load_workbook(p)
    wb["Project"]["B1"] = "Sample Plant"
    wb["Project"]["B2"] = code
    wb["Project"]["B3"] = "M. Mohanty"
    ws = wb["Cables"]
    for r, row in enumerate(rows, start=2):
        for c, col in enumerate(COLUMNS, start=1):
            ws.cell(r, c, row.get(col))
    wb.save(p)
    return p


def test_template_has_both_sheets_and_header(tmp_path):
    p = tmp_path / "t.xlsx"
    write_template(p)
    wb = load_workbook(p)
    assert wb.sheetnames == ["Project", "Cables", "Guide"]
    assert tuple(c.value for c in wb["Cables"][1]) == COLUMNS


def test_reads_project_and_rows(tmp_path):
    project, cables = read_input(filled(tmp_path, [ROW]))
    assert project == Project("Sample Plant", "IEC", "M. Mohanty")
    assert cables[0].design_current_a == 30 and cables[0].device_rating_a is None


def test_design_current_from_kw_when_blank(tmp_path):
    row = dict(ROW, design_current_a=None, load_kw=22)
    _, cables = read_input(filled(tmp_path, [row]))
    assert cables[0].design_current_a == pytest.approx(39.6928, rel=1e-4)


def test_every_problem_named_by_row(tmp_path):
    rows = [dict(ROW, length_m="fifty"), dict(ROW, tag="C2", design_current_a=None, load_kw=None)]
    with pytest.raises(InputError) as exc:
        read_input(filled(tmp_path, rows))
    assert "Row 2 (C1): field 'length_m' must be a number" in exc.value.problems
    assert "Row 3 (C2): give design_current_a or load_kw" in exc.value.problems


def test_bad_code_named(tmp_path):
    with pytest.raises(InputError, match="Project: field 'code' must be one of IEC, IS"):
        read_input(filled(tmp_path, [ROW], code="BS"))


def test_empty_rows_skipped(tmp_path):
    _, cables = read_input(filled(tmp_path, [ROW, {}]))
    assert len(cables) == 1


def test_soil_resistivity_optional_and_read(tmp_path):
    _, cables = read_input(filled(tmp_path, [ROW, dict(ROW, tag="C2", soil_resistivity_kmw=1.5)]))
    assert cables[0].soil_resistivity_kmw == 2.5
    assert cables[1].soil_resistivity_kmw == 1.5


def test_guide_explains_every_column_in_order(tmp_path):
    p = tmp_path / "t.xlsx"
    write_template(p)
    rows = list(load_workbook(p)["Guide"].iter_rows(min_row=2, max_row=len(COLUMNS) + 1,
                                                    values_only=True))
    assert tuple(r[0] for r in rows) == COLUMNS == tuple(g[0] for g in GUIDE)
    assert all(r[2] for r in rows)


def test_guide_describes_every_method(tmp_path):
    p = tmp_path / "t.xlsx"
    write_template(p)
    rows = list(load_workbook(p)["Guide"].iter_rows(min_row=len(COLUMNS) + 4, values_only=True))
    assert tuple(r[0] for r in rows) == METHODS == tuple(m[0] for m in METHOD_GUIDE)
    assert all(r[2] and "Grouping:" in r[2] for r in rows)


def test_cable_dropdown_offers_every_construction(tmp_path):
    p = tmp_path / "t.xlsx"
    write_template(p)
    ws = load_workbook(p)["Cables"]
    formulas = [dv.formula1 for dv in ws.data_validations.dataValidation]
    assert '"2 x 1C,3 x 1C,4 x 1C,2C,3C,3.5C,4C"' in formulas


def test_reader_keeps_the_sheet_row(tmp_path):
    _, cables = read_input(filled(tmp_path, [ROW, dict(ROW, tag="C2")]))
    assert [c.row for c in cables] == [2, 3]
