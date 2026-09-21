import re
from docx import Document
from cablecalc.core import size_cable
from cablecalc.io_excel import Project
from cablecalc.report import CHECK_LABELS, LIMITATIONS, write_report
from tests.test_models import make

BANNED = re.compile(r"[→·•\U0001F300-\U0001FAFF☀-➿]")


def all_text(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def build(tmp_path, fixture_data):
    results = [size_cable(make(), fixture_data),
               size_cable(make(tag="C2", length_m=400), fixture_data),
               size_cable(make(tag="C3", ambient_c=35), fixture_data)]
    p = tmp_path / "r.docx"
    write_report(p, Project("Sample Plant", "FIXTURE", "M. Mohanty"), results, today="2026-09-19")
    return all_text(p)


def test_cover_and_blank_checker(tmp_path, fixture_data):
    t = build(tmp_path, fixture_data)
    assert "Sample Plant" in t and "2026-09-19" in t and "FIXTURE" in t
    assert "Checked by: ____________________" in t


def test_code_named_in_full_and_pages_numbered(tmp_path, fixture_data):
    from cablecalc.report import CODE_TITLES
    p = tmp_path / "r.docx"
    write_report(p, Project("Sample Plant", "IS", "M. Mohanty"),
                 [size_cable(make(), fixture_data)], today="2026-09-21")
    d = Document(p)
    assert CODE_TITLES["IS"] in all_text(p)
    footer_xml = d.sections[0].footer._element.xml
    assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml


def test_every_check_with_formula_and_result(tmp_path, fixture_data):
    t = build(tmp_path, fixture_data)
    assert "Iz = It x Ca x Cg >= max(Ib, In)" in t and "S >= I x sqrt(t) / k" in t
    assert "PASS" in t and "FAIL" in t
    assert "No size passes" in t


def test_checks_use_plain_labels_not_code_names(tmp_path, fixture_data):
    t = build(tmp_path, fixture_data)
    for label in CHECK_LABELS.values():
        assert label in t
    for name in ("voltage_drop", "short_circuit"):  # "current" is also an ordinary word
        assert name not in t


def test_every_source_printed(tmp_path, fixture_data):
    assert "TEST FIXTURE - not from any standard" in build(tmp_path, fixture_data)


def test_rounding_note_reaches_the_report(tmp_path, fixture_data):
    assert "35 deg C taken as 40 deg C, next higher tabulated value" in build(tmp_path, fixture_data)


def test_summary_and_limitations(tmp_path, fixture_data):
    t = build(tmp_path, fixture_data)
    assert "Summary" in t and "Short-circuit withstand" in t
    for line in LIMITATIONS:
        assert line in t


def test_no_banned_characters(tmp_path, fixture_data):
    assert not BANNED.search(build(tmp_path, fixture_data))


def test_governing_check_is_defined_in_the_report(tmp_path, fixture_data):
    from cablecalc.report import GOVERNING_NOTE
    assert GOVERNING_NOTE in build(tmp_path, fixture_data)
