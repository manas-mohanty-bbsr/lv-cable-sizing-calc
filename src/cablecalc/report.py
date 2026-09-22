"""Word out: a formal calculation sheet. Plain text only - no emoji, arrows or middots."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

from . import __version__
from .io_excel import Project
from .models import SizingResult

CHECK_LABELS = {
    "current": "Current-carrying capacity",
    "voltage_drop": "Voltage drop",
    "short_circuit": "Short-circuit withstand",
}

GOVERNING_NOTE = ("Governing check: the check that the next smaller size fails (the one failing by the "
                  "widest margin, if more than one does). A dash means the smallest size in the table "
                  "already passes.")

LIMITATIONS = (
    "Low-voltage cables up to 1.1 kV only.",
    "Only the installation methods and sizes present in the shipped tables are supported. An ambient "
    "temperature, number of grouped circuits or soil thermal resistivity between tabulated rows takes "
    "the next higher row, as stated in the Sources line; values are never interpolated.",
    "For buried cables (methods D1 and D2) the ambient temperature is the ground temperature, and "
    "grouping assumes the cables or ducts are touching.",
    "Group reduction factors assume a group of similar, equally loaded cables. Groups containing "
    "different sizes are not covered; for those, IS 732 clause S-5.1 gives F = 1/sqrt(n).",
    "Overload protection: condition (1) of IS 732 clause 4.4.4.1, Ib <= In <= Iz, is checked. Condition "
    "(2), I2 <= 1.45 x Iz, is not calculated. It is met by a device whose conventional tripping current "
    "I2 is not more than 1.45 x In, such as a circuit-breaker to IEC 60898-1; for any other device the "
    "engineer must check condition (2).",
    "Short-circuit withstand is checked as S >= I x sqrt(t) / k (IS 732 clause 4.4.5.5.2, equation 3). "
    "The same clause requires k^2 x S^2 to exceed the let-through energy I^2t quoted by the device "
    "manufacturer where the device operates in less than 0.1 s or is current-limiting. That check is "
    "not calculated; the engineer must make it from the manufacturer's data.",
    "The fault current and fault time entered must be the worst pair for a fault anywhere along the "
    "cable (IS 732 clause 4.4.5.5.2). A fault at the far end is smaller but may take longer to clear, so "
    "it can be the more onerous case. The tool checks only the pair entered.",
    "Harmonic derating is not calculated.",
    "This is a sample tool, not a design service. Results must be checked by a qualified engineer.",
)


CODE_TITLES = {
    "IEC": "IEC 60364-5-52:2009 (values as reproduced in IS 732:2019)",
    "IS": "IS 732:2019",
}

LABEL_WIDTH, VALUE_WIDTH = Cm(3.5), Cm(12.5)


def _table(doc, rows, widths=None):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Table Grid"
    if widths:
        t.autofit = False
    for row in rows:
        cells = t.add_row().cells
        for i, (c, v) in enumerate(zip(cells, row)):
            c.text = str(v)
            if widths:
                c.width = widths[i]
    return t


def _num(x):
    """Two decimal places, trailing zeros dropped: 50 -> 50, 13.75 -> 13.75, 0.03399 -> 0.03."""
    return f"{x:.2f}".rstrip("0").rstrip(".")


def _keep_together(table):
    """Stop a table splitting across a page break."""
    rows = table.rows
    for i, row in enumerate(rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        tr_pr.append(cant)
        if i < len(rows) - 1:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.keep_with_next = True


def _field(run, instr):
    for kind, text in (("begin", None), (None, instr), ("separate", None), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _page_numbers(doc):
    para = doc.sections[0].footer.paragraphs[0]
    para.add_run("Page ")
    _field(para.add_run(), "PAGE")
    para.add_run(" of ")
    _field(para.add_run(), "NUMPAGES")


def write_report(path: Path, project: Project, results: list[SizingResult], today: str) -> None:
    doc = Document()
    doc.add_heading("LV Cable Sizing Calculation", level=0)
    _page_numbers(doc)
    _table(doc, [("Project", project.name), ("Date", today),
                 ("Code", CODE_TITLES.get(project.code, project.code)),
                 ("Tool version", f"cablecalc {__version__}"),
                 ("Prepared by", project.prepared_by)], widths=(LABEL_WIDTH, VALUE_WIDTH))
    doc.add_paragraph("Checked by: ____________________")

    for res in results:
        doc.add_heading(f"Cable {res.tag}", level=1)
        doc.add_paragraph(res.message).paragraph_format.keep_with_next = True
        for c in res.checks:
            doc.add_heading(CHECK_LABELS[c.name], level=2)
            t = _table(doc, [
                ("Formula", c.formula),
                ("Values", c.detail),
                ("Required", f"{_num(c.required)} {c.unit}"),
                ("Actual", f"{_num(c.actual)} {c.unit}"),
                ("Result", "PASS" if c.passed else "FAIL"),
                ("Sources", "\n".join(c.sources)),
            ], widths=(LABEL_WIDTH, VALUE_WIDTH))
            _keep_together(t)

    doc.add_heading("Summary", level=1)
    summary = _table(doc, [("Cable", "Size (mm2)", "Governing check", "Result")] + [
        (r.tag, f"{r.size_mm2:g}" if r.size_mm2 else "-",
         CHECK_LABELS[r.governing] if r.governing else "-",
         "PASS" if r.passed else "FAIL") for r in results])
    _keep_together(summary)
    for cell in summary.rows[-1].cells:  # keep the governing note on the table's page
        for para in cell.paragraphs:
            para.paragraph_format.keep_with_next = True
    doc.add_paragraph(GOVERNING_NOTE)

    doc.add_heading("Limitations", level=1)
    for line in LIMITATIONS:
        doc.add_paragraph(line)
    doc.save(path)
