"""Excel in: a blank template, and a reader that reports every problem at once."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .codes import SUPPORTED_CODES
from .models import CableInput, InputError, design_current_from_kw, validate_input

COLUMNS = ("tag", "phase", "voltage_v", "design_current_a", "load_kw", "power_factor",
           "length_m", "conductor", "insulation", "method", "cores", "ambient_c",
           "grouped_circuits", "vd_limit_pct", "fault_current_ka", "fault_time_s",
           "device_rating_a", "soil_resistivity_kmw")
_TEXT = {"tag", "phase", "conductor", "insulation", "method"}
_INT = {"cores", "grouped_circuits"}
_OPTIONAL = {"design_current_a", "load_kw", "device_rating_a", "soil_resistivity_kmw"}


@dataclass(frozen=True)
class Project:
    name: str
    code: str
    prepared_by: str


METHODS = ("A1", "A2", "B1", "B2", "C", "E", "D1", "D2")

# (column, unit, what to enter) - written to the Guide sheet of the template
GUIDE = (
    ("tag", "-", "Cable reference, e.g. F-01. Required."),
    ("phase", "-", "1ph or 3ph."),
    ("voltage_v", "V", "Line voltage for 3ph (e.g. 415), phase voltage for 1ph (e.g. 230)."),
    ("design_current_a", "A", "Design current Ib. Leave blank to calculate it from load_kw."),
    ("load_kw", "kW", "Used only when design_current_a is blank."),
    ("power_factor", "-", "Between 0 and 1, e.g. 0.85."),
    ("length_m", "m", "Route length, one way."),
    ("conductor", "-", "Cu or Al."),
    ("insulation", "-", "PVC or XLPE."),
    ("method", "-", "Installation method: " + ", ".join(METHODS) + ". Each is described in the "
     "Installation methods list below."),
    ("cores", "-", "LOADED conductors: 2 for 1ph, 3 for 3ph (a 4-core 3ph cable is entered as 3)."),
    ("ambient_c", "deg C", "Air temperature; for D1 and D2 the ground temperature. Between table rows, "
     "the next higher row is used."),
    ("grouped_circuits", "-", "Number of circuits in the group, including this one. 1 if alone."),
    ("vd_limit_pct", "%", "Permitted voltage drop, e.g. 5."),
    ("fault_current_ka", "kA", "Prospective fault current at the cable."),
    ("fault_time_s", "s", "Disconnection time of the protective device for that fault."),
    ("device_rating_a", "A", "Rating In of the overload device. Leave blank if there is none."),
    ("soil_resistivity_kmw", "K.m/W", "D1 and D2 only. Leave blank for 2.5."),
)

# (method, where the cable runs and the grouping factors applied) - written below GUIDE.
# From IS 732:2019 Table 19, Table 20 and clause S-6.1; grouping rows from Tables 36 to 38. The codes
# mean the same in IEC 60364-5-52:2009 (Tables A.52.3 and B.52.1).
METHOD_GUIDE = (
    ("A1", "Single-core cables or insulated conductors in conduit in a thermally insulated wall. "
     "Grouping: bunched or enclosed (Table 36 item 1)."),
    ("A2", "Multi-core cable in conduit in a thermally insulated wall. "
     "Grouping: bunched or enclosed (Table 36 item 1)."),
    ("B1", "Single-core cables or insulated conductors in conduit or trunking on a wall, less than "
     "0.3 x conduit diameter from it. Grouping: bunched or enclosed (Table 36 item 1)."),
    ("B2", "Multi-core cable in conduit on a wall, less than 0.3 x conduit diameter from it. "
     "Grouping: bunched or enclosed (Table 36 item 1)."),
    ("C", "Single-core or multi-core cable fixed on a wall, less than 0.3 x cable diameter from it; "
     "also used for cable on unperforated tray and cable direct in masonry. "
     "Grouping: single layer on wall, floor or unperforated tray (Table 36 item 2)."),
    ("E", "Multi-core cable in free air, at least 0.3 x cable diameter clear of any surface, for "
     "example on perforated tray, brackets or ladder. "
     "Grouping: single layer on perforated tray (Table 36 item 4)."),
    ("D1", "Multi-core cable in ducts in the ground (reference: 100 mm duct, 0.7 m deep). "
     "Grouping: ducts touching (Table 38)."),
    ("D2", "Cable laid direct in the ground (reference: 0.7 m deep). "
     "Grouping: cables touching (Table 37)."),
)


def _dropdown(ws, choices, cells):
    dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"')
    ws.add_data_validation(dv)
    dv.add(cells)


def write_template(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Project"
    for r, label in enumerate(("Project name", "Code (IEC or IS)", "Prepared by"), start=1):
        ws.cell(r, 1, label)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 40
    _dropdown(ws, SUPPORTED_CODES, "B2")

    cables = wb.create_sheet("Cables")
    cables.append(list(COLUMNS))
    for i, col in enumerate(COLUMNS, start=1):
        cables.column_dimensions[get_column_letter(i)].width = max(10, len(col) + 2)
    for col, choices in (("phase", ("1ph", "3ph")), ("conductor", ("Cu", "Al")),
                         ("insulation", ("PVC", "XLPE")), ("method", METHODS)):
        letter = get_column_letter(COLUMNS.index(col) + 1)
        _dropdown(cables, choices, f"{letter}2:{letter}500")

    write_guide(wb)
    wb.save(path)


def write_guide(wb: Workbook) -> None:
    """Add the Guide sheet: every column, then every installation method."""
    guide = wb.create_sheet("Guide")
    guide.append(["Column", "Unit", "What to enter"])
    for row in GUIDE:
        guide.append(list(row))
    guide.append([])
    guide.append(["Installation methods", "", "Where the cable runs, and the grouping factors used"])
    for method, text in METHOD_GUIDE:
        guide.append([method, "", text])
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 8
    guide.column_dimensions["C"].width = 100


def read_input(path: Path) -> tuple[Project, list[CableInput]]:
    wb = load_workbook(path, data_only=True)
    ps = wb["Project"]
    problems: list[str] = []
    code = str(ps["B2"].value or "").strip().upper()
    if code not in SUPPORTED_CODES:
        problems.append(f"Project: field 'code' must be one of {', '.join(SUPPORTED_CODES)}")
    project = Project(str(ps["B1"].value or ""), code, str(ps["B3"].value or ""))

    cables: list[CableInput] = []
    for r, values in enumerate(wb["Cables"].iter_rows(min_row=2, values_only=True), start=2):
        raw = dict(zip(COLUMNS, values))
        if all(v in (None, "") for v in raw.values()):
            continue
        tag = str(raw["tag"] or f"row{r}")
        where = f"Row {r} ({tag})"
        row: dict = {}
        for col in COLUMNS:
            v = raw[col]
            if col in _TEXT:
                row[col] = str(v).strip() if v is not None else ""
                if not row[col]:
                    problems.append(f"{where}: field '{col}' is empty")
            elif v in (None, ""):
                row[col] = None
                if col not in _OPTIONAL:
                    problems.append(f"{where}: field '{col}' is empty")
            else:
                try:
                    row[col] = int(v) if col in _INT else float(v)
                except (TypeError, ValueError):
                    problems.append(f"{where}: field '{col}' must be a number")
                    row[col] = None
        if row.get("design_current_a") is None:
            if row.get("load_kw") is None:
                problems.append(f"{where}: give design_current_a or load_kw")
                continue
            if None in (row.get("voltage_v"), row.get("power_factor")):
                continue
            row["design_current_a"] = design_current_from_kw(
                row["load_kw"], row["voltage_v"], row["power_factor"], row["phase"])
        if any(row[c] is None for c in COLUMNS if c not in _OPTIONAL):
            continue
        row.pop("load_kw")
        if row["soil_resistivity_kmw"] is None:
            row.pop("soil_resistivity_kmw")
        inp = CableInput(**row)
        try:
            validate_input(inp)
        except InputError as e:
            problems.extend(f"Row {r} - {p}" for p in e.problems)
            continue
        cables.append(inp)
    if problems:
        raise InputError(problems)
    return project, cables
